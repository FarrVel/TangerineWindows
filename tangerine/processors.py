"""
Fungsi-fungsi pemrosesan file 100% OFFLINE untuk Tangerine for Windows.

Setiap fungsi menerima path sumber (dan argumen opsional) lalu mengembalikan
objek `Result` yang berisi daftar file output, pesan, dan (opsional) teks
yang akan disalin ke clipboard.
"""
from __future__ import annotations

import html
import io
import os
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

# --------------------------------------------------------------------------- #
#  Util umum
# --------------------------------------------------------------------------- #

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff",
              ".heic", ".heif", ".avif", ".ico", ".jfif"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
TEXT_EXTS = {".txt", ".md", ".markdown", ".html", ".htm", ".csv", ".log", ".json", ".xml"}
OFFICE_EXTS = {".docx", ".xlsx", ".pptx"}


@dataclass
class Result:
    outputs: list[str] = field(default_factory=list)
    message: str = ""
    clipboard: str | None = None


class ProcessingError(Exception):
    """Error yang pesannya aman ditampilkan ke pengguna."""


def unique_path(src: str | Path, suffix: str, ext: str, extra: str = "") -> Path:
    """dokumen.pdf -> dokumen_converted.docx (tidak menimpa file yang sudah ada)."""
    p = Path(src)
    stem = p.stem if p.is_file() else p.name
    if not ext.startswith(".") and ext:
        ext = "." + ext
    base = p.parent / f"{stem}_{suffix}{extra}{ext}"
    i = 1
    while base.exists():
        base = p.parent / f"{stem}_{suffix}{extra} ({i}){ext}"
        i += 1
    return base


def human_size(n: float) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} GB"


def _size_report(src: Path, out: Path) -> str:
    a, b = src.stat().st_size, out.stat().st_size
    pct = (1 - b / a) * 100 if a else 0
    return f"{human_size(a)} → {human_size(b)} ({pct:+.0f}% lebih kecil)" if pct >= 0 \
        else f"{human_size(a)} → {human_size(b)} (file sudah optimal)"


# --------------------------------------------------------------------------- #
#  Pillow helpers
# --------------------------------------------------------------------------- #

_HEIF_READY = False


def _pil():
    """Import Pillow + daftarkan plugin HEIC/AVIF bila tersedia."""
    global _HEIF_READY
    from PIL import Image
    if not _HEIF_READY:
        try:
            import pillow_heif
            pillow_heif.register_heif_opener()
            if hasattr(pillow_heif, "register_avif_opener"):
                try:
                    pillow_heif.register_avif_opener()
                except Exception:
                    pass
        except Exception:
            pass
        try:  # plugin AVIF terpisah (opsional)
            import pillow_avif  # noqa: F401
        except Exception:
            pass
        _HEIF_READY = True
    return Image


def open_image(path: str | Path):
    Image = _pil()
    from PIL import ImageOps
    img = Image.open(path)
    try:
        img.seek(0)
    except Exception:
        pass
    img.load()
    try:
        img = ImageOps.exif_transpose(img)
    except Exception:
        pass
    return img


def _flatten(img, color=(255, 255, 255)):
    """Konversi ke RGB, transparansi digabung dengan warna latar."""
    Image = _pil()
    if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
        rgba = img.convert("RGBA")
        bg = Image.new("RGB", rgba.size, color)
        bg.paste(rgba, mask=rgba.split()[-1])
        return bg
    return img.convert("RGB")


def _avif_supported() -> bool:
    Image = _pil()
    from PIL import features
    try:
        if features.check("avif"):
            return True
    except Exception:
        pass
    Image.init()
    return "AVIF" in Image.SAVE


def save_image(img, out: Path, fmt: str, quality: int | None = None, optimize=True):
    """Simpan PIL Image ke format tertentu dengan penanganan mode warna."""
    Image = _pil()
    fmt = fmt.upper()
    if fmt in ("JPG", "JPEG", "JFIF"):
        _flatten(img).save(out, "JPEG", quality=quality or 92, optimize=optimize,
                           progressive=True, subsampling="4:2:0" if (quality or 92) < 90 else 0)
    elif fmt == "PNG":
        if img.mode not in ("RGB", "RGBA", "L", "LA", "P", "1", "I", "I;16"):
            img = img.convert("RGBA")
        img.save(out, "PNG", optimize=optimize)
    elif fmt == "WEBP":
        im = img if img.mode in ("RGB", "RGBA") else img.convert("RGBA")
        im.save(out, "WEBP", quality=quality or 90, method=6)
    elif fmt == "BMP":
        _flatten(img).save(out, "BMP")
    elif fmt in ("TIFF", "TIF"):
        im = img if img.mode in ("RGB", "RGBA", "L", "CMYK") else img.convert("RGBA")
        im.save(out, "TIFF", compression="tiff_lzw")
    elif fmt == "GIF":
        img.convert("RGBA").save(out, "GIF")
    elif fmt == "ICO":
        im = img.convert("RGBA")
        sizes = [s for s in [(16, 16), (24, 24), (32, 32), (48, 48), (64, 64),
                             (128, 128), (256, 256)] if s[0] <= max(im.size)]
        im.save(out, "ICO", sizes=sizes or [(16, 16)])
    elif fmt in ("HEIC", "HEIF"):
        try:
            import pillow_heif  # noqa: F401
        except Exception:
            raise ProcessingError("Dukungan HEIC butuh library 'pillow-heif'.")
        im = img if img.mode in ("RGB", "RGBA") else img.convert("RGBA")
        im.save(out, "HEIF", quality=quality or 90)
    elif fmt == "AVIF":
        if not _avif_supported():
            raise ProcessingError("Encoder AVIF tidak tersedia. Jalankan: pip install -U pillow "
                                  "atau pip install pillow-avif-plugin")
        im = img if img.mode in ("RGB", "RGBA") else img.convert("RGBA")
        im.save(out, "AVIF", quality=quality or 80)
    elif fmt == "PDF":
        _flatten(img).save(out, "PDF", resolution=150.0)
    else:
        img.save(out, fmt)


# --------------------------------------------------------------------------- #
#  QR Code
# --------------------------------------------------------------------------- #

def _decode_qr_pil(img) -> list[str]:
    """Decode QR/barcode dari PIL image. Mencoba zxing-cpp lalu OpenCV."""
    found: list[str] = []
    rgb = _flatten(img)
    try:
        import zxingcpp
        for r in zxingcpp.read_barcodes(rgb):
            if r.text and r.text not in found:
                found.append(r.text)
        if found:
            return found
    except ImportError:
        pass
    except Exception:
        pass
    try:
        import cv2
        import numpy as np
        arr = cv2.cvtColor(np.array(rgb), cv2.COLOR_RGB2BGR)
        det = cv2.QRCodeDetector()
        try:
            ok, infos, _pts, _ = det.detectAndDecodeMulti(arr)
            if ok:
                for t in infos:
                    if t and t not in found:
                        found.append(t)
        except Exception:
            pass
        if not found:
            t, _pts, _ = det.detectAndDecode(arr)
            if t:
                found.append(t)
    except ImportError:
        if not found:
            raise ProcessingError("Pembaca QR butuh 'zxing-cpp' atau 'opencv-python-headless'.")
    return found


def _qr_result(src: Path, codes: list[str]) -> Result:
    if not codes:
        return Result(message=f"Tidak ada QR code yang terdeteksi di {src.name}.")
    text = "\n".join(codes)
    out = unique_path(src, "qr", ".txt")
    out.write_text(text, encoding="utf-8")
    preview = text if len(text) < 120 else text[:117] + "..."
    return Result([str(out)], f"{len(codes)} QR ditemukan (disalin ke clipboard):\n{preview}", text)


def image_read_qr(path: str) -> Result:
    src = Path(path)
    return _qr_result(src, _decode_qr_pil(open_image(src)))


def pdf_scan_qr(path: str) -> Result:
    fitz = _fitz()
    Image = _pil()
    src = Path(path)
    codes: list[str] = []
    with fitz.open(src) as doc:
        for page in doc:
            pix = page.get_pixmap(dpi=200, alpha=False)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            for c in _decode_qr_pil(img):
                if c not in codes:
                    codes.append(c)
    return _qr_result(src, codes)


# --------------------------------------------------------------------------- #
#  Gambar
# --------------------------------------------------------------------------- #

def image_convert(path: str, fmt: str) -> Result:
    src = Path(path)
    img = open_image(src)
    ext = {"JPEG": ".jpg", "JPG": ".jpg", "TIFF": ".tiff", "HEIF": ".heic"}.get(fmt.upper(),
                                                                            "." + fmt.lower())
    out = unique_path(src, "converted", ext)
    save_image(img, out, fmt)
    return Result([str(out)], f"{src.name} → {out.name}")


def images_to_pdf(paths: list[str]) -> Result:
    """Satu gambar -> satu PDF; banyak gambar -> digabung jadi satu PDF."""
    imgs = [_flatten(open_image(p)) for p in paths]
    out = unique_path(paths[0], "converted" if len(paths) == 1 else "merged", ".pdf")
    imgs[0].save(out, "PDF", resolution=150.0, save_all=True, append_images=imgs[1:])
    return Result([str(out)], f"{len(paths)} gambar → {out.name}")


def image_to_docx(path: str) -> Result:
    from docx import Document
    from docx.shared import Inches
    src = Path(path)
    img = open_image(src)
    buf = io.BytesIO()
    (img.convert("RGBA") if img.mode in ("RGBA", "LA", "P") else _flatten(img)).save(buf, "PNG")
    buf.seek(0)
    doc = Document()
    w, h = img.size
    max_w, max_h = 6.3, 8.8
    if w / h > max_w / max_h:
        doc.add_picture(buf, width=Inches(max_w))
    else:
        doc.add_picture(buf, height=Inches(max_h))
    out = unique_path(src, "converted", ".docx")
    doc.save(out)
    return Result([str(out)], f"{src.name} → {out.name}")


def image_compress(path: str) -> Result:
    src = Path(path)
    img = open_image(src)
    fmt = (Image_format(src) or "JPEG").upper()
    # perkecil gambar yang sangat besar
    max_side = 3840
    if max(img.size) > max_side:
        img.thumbnail((max_side, max_side))
    if fmt in ("JPEG", "MPO"):
        out = unique_path(src, "compressed", src.suffix or ".jpg")
        save_image(img, out, "JPEG", quality=72)
    elif fmt == "PNG":
        out = unique_path(src, "compressed", ".png")
        im = img.convert("RGBA") if img.mode not in ("RGB", "RGBA") else img
        try:
            q = im.quantize(colors=256, method=2 if im.mode == "RGBA" else 0, dither=1)
        except Exception:
            q = im.convert("P", palette=1, colors=256)
        q.save(out, "PNG", optimize=True)
    elif fmt == "WEBP":
        out = unique_path(src, "compressed", ".webp")
        save_image(img, out, "WEBP", quality=70)
    elif fmt in ("HEIF", "HEIC", "AVIF"):
        out = unique_path(src, "compressed", src.suffix)
        save_image(img, out, "HEIC" if fmt != "AVIF" else "AVIF", quality=60)
    else:
        out = unique_path(src, "compressed", ".jpg")
        save_image(img, out, "JPEG", quality=72)
    if out.stat().st_size >= src.stat().st_size and fmt != "BMP":
        out.unlink(missing_ok=True)
        return Result([], f"{src.name} sudah teroptimasi, tidak bisa dikompres lebih kecil.")
    return Result([str(out)], f"{out.name}\n{_size_report(src, out)}")


def Image_format(src: Path) -> str | None:
    Image = _pil()
    try:
        with Image.open(src) as im:
            return im.format
    except Exception:
        return None


def image_strip_metadata(path: str) -> Result:
    src = Path(path)
    img = open_image(src)  # orientasi EXIF sudah diterapkan ke piksel
    fmt = (Image_format(src) or "PNG").upper()
    icc = img.info.get("icc_profile")
    img.info = {}
    if icc:
        img.info["icc_profile"] = icc
    ext = src.suffix if src.suffix else ".png"
    out = unique_path(src, "clean", ext)
    fmt_map = {"MPO": "JPEG", "JPEG": "JPEG", "PNG": "PNG", "WEBP": "WEBP", "TIFF": "TIFF",
               "BMP": "BMP", "GIF": "GIF", "HEIF": "HEIC", "AVIF": "AVIF", "ICO": "ICO"}
    save_image(img, out, fmt_map.get(fmt, "PNG"), quality=95)
    return Result([str(out)], f"Metadata (EXIF/GPS/kamera) dihapus → {out.name}")


def image_add_background(path: str, color=(255, 255, 255)) -> Result:
    Image = _pil()
    src = Path(path)
    img = open_image(src).convert("RGBA")
    bg = Image.new("RGBA", img.size, tuple(color) + (255,))
    bg.alpha_composite(img)
    ext = src.suffix.lower() if src.suffix.lower() in (".png", ".jpg", ".jpeg", ".webp", ".bmp") else ".png"
    out = unique_path(src, "bg", ext)
    save_image(bg.convert("RGB"), out, {".jpg": "JPEG", ".jpeg": "JPEG"}.get(ext, ext[1:].upper()))
    return Result([str(out)], f"Latar belakang ditambahkan → {out.name}")


# --------------------------------------------------------------------------- #
#  PDF
# --------------------------------------------------------------------------- #

def _fitz():
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore
    return fitz


def pdf_to_docx(path: str) -> Result:
    src = Path(path)
    out = unique_path(src, "converted", ".docx")
    try:
        from pdf2docx import Converter
    except ImportError:
        raise ProcessingError("Library 'pdf2docx' belum terpasang.")
    import logging
    logging.getLogger().setLevel(logging.ERROR)
    cv = Converter(str(src))
    try:
        cv.convert(str(out), multi_processing=False)
    finally:
        cv.close()
    return Result([str(out)], f"{src.name} → {out.name}")


def pdf_to_txt(path: str) -> Result:
    fitz = _fitz()
    src = Path(path)
    parts = []
    with fitz.open(src) as doc:
        for i, page in enumerate(doc, 1):
            parts.append(page.get_text("text", sort=True).strip())
    text = "\n\n\f".join(parts)
    out = unique_path(src, "converted", ".txt")
    out.write_text(text, encoding="utf-8")
    msg = f"{src.name} → {out.name}"
    if not text.strip():
        msg += "\n(PDF tampaknya hasil scan — tidak ada teks yang bisa diambil)"
    return Result([str(out)], msg)


def pdf_to_images(path: str, fmt: str = "PNG", dpi: int = 150) -> Result:
    fitz = _fitz()
    Image = _pil()
    src = Path(path)
    outs: list[str] = []
    ext = ".jpg" if fmt.upper() in ("JPG", "JPEG") else "." + fmt.lower()
    with fitz.open(src) as doc:
        n = doc.page_count
        width = max(2, len(str(n)))
        for i, page in enumerate(doc, 1):
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            extra = "" if n == 1 else f"_page{str(i).zfill(width)}"
            out = unique_path(src, "converted", ext, extra)
            save_image(img, out, fmt, quality=90)
            outs.append(str(out))
    return Result(outs, f"{src.name} → {len(outs)} gambar {fmt.upper()}")


def pdf_to_epub(path: str) -> Result:
    fitz = _fitz()
    src = Path(path)
    chapters = []
    with fitz.open(src) as doc:
        title = (doc.metadata or {}).get("title") or src.stem
        author = (doc.metadata or {}).get("author") or "Tidak diketahui"
        for i, page in enumerate(doc, 1):
            paras = []
            for b in page.get_text("blocks", sort=True):
                if b[6] == 0:  # blok teks
                    t = " ".join(b[4].split())
                    if t:
                        paras.append(f"<p>{html.escape(t)}</p>")
            chapters.append((f"Halaman {i}", "\n".join(paras) or "<p></p>"))
    out = unique_path(src, "converted", ".epub")
    build_epub(out, title, author, chapters)
    return Result([str(out)], f"{src.name} → {out.name}")


def pdf_compress(path: str) -> Result:
    fitz = _fitz()
    src = Path(path)
    out = unique_path(src, "compressed", ".pdf")
    with fitz.open(src) as doc:
        if hasattr(doc, "rewrite_images"):
            try:
                doc.rewrite_images(dpi_threshold=160, dpi_target=120, quality=65)
            except Exception:
                pass
        try:
            doc.scrub(metadata=False, xml_metadata=False, thumbnails=True,
                      reset_fields=False, reset_responses=False, javascript=False,
                      attached_files=False, clean_pages=True, embedded_files=False,
                      hidden_text=False, redactions=False, remove_links=False)
        except Exception:
            pass
        try:
            doc.subset_fonts()
        except Exception:
            pass
        kwargs = dict(garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
        try:
            doc.save(out, use_objstms=1, **kwargs)
        except TypeError:
            doc.save(out, **kwargs)
    if out.stat().st_size >= src.stat().st_size:
        out.unlink(missing_ok=True)
        return Result([], f"{src.name} sudah teroptimasi, tidak bisa dikompres lebih kecil.")
    return Result([str(out)], f"{out.name}\n{_size_report(src, out)}")


def pdf_strip_metadata(path: str) -> Result:
    fitz = _fitz()
    src = Path(path)
    out = unique_path(src, "clean", ".pdf")
    with fitz.open(src) as doc:
        doc.set_metadata({})
        try:
            doc.del_xml_metadata()
        except Exception:
            pass
        doc.save(out, garbage=3, deflate=True)
    return Result([str(out)], f"Metadata PDF dihapus → {out.name}")


def pdfs_merge(paths: list[str]) -> Result:
    fitz = _fitz()
    out = unique_path(paths[0], "merged", ".pdf")
    with fitz.open() as dst:
        for p in paths:
            with fitz.open(p) as s:
                dst.insert_pdf(s)
        dst.save(out, garbage=3, deflate=True)
    return Result([str(out)], f"{len(paths)} PDF digabung → {out.name}")


# --------------------------------------------------------------------------- #
#  HTML -> PDF (MuPDF Story) dipakai untuk DOCX/TXT/MD -> PDF tanpa Word
# --------------------------------------------------------------------------- #

_CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.45; color: #111; }
h1 { font-size: 20pt; } h2 { font-size: 16pt; } h3 { font-size: 13pt; }
table { border-collapse: collapse; } td, th { border: 1px solid #999; padding: 3px; }
pre, code { font-family: monospace; font-size: 9.5pt; }
img { max-width: 100%; }
"""


def html_to_pdf(html_str: str, out: Path, base_dir: str | None = None):
    fitz = _fitz()
    archive = fitz.Archive(base_dir) if base_dir else None
    story = fitz.Story(html=html_str, user_css=_CSS, archive=archive)
    writer = fitz.DocumentWriter(str(out))
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (54, 54, -54, -54)
    more = 1
    pages = 0
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
        pages += 1
        if pages > 5000:
            break
    writer.close()


def _text_to_html(text: str, markdown: bool = False) -> str:
    if markdown:
        try:
            import markdown as md
            return md.markdown(text, extensions=["tables", "fenced_code"])
        except ImportError:
            pass
    paras = [f"<p>{html.escape(p).replace(chr(10), '<br/>')}</p>" for p in text.split("\n\n")]
    return "\n".join(paras)


def _read_text(path: Path) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_bytes().decode("utf-8", "replace")


# --------------------------------------------------------------------------- #
#  DOCX
# --------------------------------------------------------------------------- #

def _find_soffice() -> str | None:
    cands = [shutil.which("soffice"), shutil.which("soffice.exe")]
    for base in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                 os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
        cands.append(os.path.join(base, "LibreOffice", "program", "soffice.exe"))
    for c in cands:
        if c and os.path.isfile(c):
            return c
    return None


def _docx_pdf_via_word(src: Path, out: Path) -> bool:
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return False
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(src), ReadOnly=True, AddToRecentFiles=False,
                                  ConfirmConversions=False, Visible=False)
        doc.ExportAsFixedFormat(str(out), 17)  # 17 = wdExportFormatPDF
        doc.Close(False)
        return out.exists()
    except Exception:
        return False
    finally:
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


def _docx_pdf_via_libreoffice(src: Path, out: Path) -> bool:
    soffice = _find_soffice()
    if not soffice:
        return False
    tmp = Path(tempfile.mkdtemp(prefix="tangerine_"))
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(src)],
                       timeout=180, capture_output=True,
                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        res = tmp / (src.stem + ".pdf")
        if res.exists():
            shutil.move(str(res), str(out))
            return True
        return False
    except Exception:
        return False
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _docx_to_html(src: Path) -> str:
    try:
        import mammoth
    except ImportError:
        raise ProcessingError("Library 'mammoth' belum terpasang.")
    with open(src, "rb") as f:
        res = mammoth.convert_to_html(f)
    return res.value


def docx_to_pdf(path: str) -> Result:
    src = Path(path)
    out = unique_path(src, "converted", ".pdf")
    engine = None
    if _docx_pdf_via_word(src, out):
        engine = "Microsoft Word"
    elif _docx_pdf_via_libreoffice(src, out):
        engine = "LibreOffice"
    else:
        html_to_pdf(_docx_to_html(src), out)
        engine = "mesin internal"
    return Result([str(out)], f"{src.name} → {out.name}\n(via {engine})")


def docx_to_txt(path: str) -> Result:
    from docx import Document
    src = Path(path)
    doc = Document(str(src))
    lines = []
    body = doc.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(child, doc).text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                lines.append("\t".join(c.text.strip() for c in row.cells))
            lines.append("")
    out = unique_path(src, "converted", ".txt")
    out.write_text("\n".join(lines), encoding="utf-8")
    return Result([str(out)], f"{src.name} → {out.name}")


def docx_to_html(path: str) -> Result:
    src = Path(path)
    body = _docx_to_html(src)
    page = (f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{html.escape(src.stem)}</title>"
            f"<style>{_CSS} body{{max-width:820px;margin:40px auto;}}</style></head>"
            f"<body>{body}</body></html>")
    out = unique_path(src, "converted", ".html")
    out.write_text(page, encoding="utf-8")
    return Result([str(out)], f"{src.name} → {out.name}")


def docx_to_epub(path: str) -> Result:
    src = Path(path)
    body = _docx_to_html(src)
    # gambar inline (data URI) dibuang agar EPUB tetap valid & ringan
    import re
    body = re.sub(r"<img[^>]*>", "", body)
    out = unique_path(src, "converted", ".epub")
    build_epub(out, src.stem, "Tidak diketahui", [(src.stem, _xhtml_safe(body))])
    return Result([str(out)], f"{src.name} → {out.name}")


def docx_strip_metadata(path: str) -> Result:
    from docx import Document
    src = Path(path)
    doc = Document(str(src))
    cp = doc.core_properties
    for attr in ("author", "category", "comments", "content_status", "identifier", "keywords",
                 "language", "last_modified_by", "subject", "title", "version"):
        try:
            setattr(cp, attr, "")
        except Exception:
            pass
    try:
        cp.revision = 1
        cp.created = datetime(2000, 1, 1)
        cp.modified = datetime(2000, 1, 1)
        cp.last_printed = datetime(2000, 1, 1)
    except Exception:
        pass
    out = unique_path(src, "clean", ".docx")
    doc.save(out)
    return Result([str(out)], f"Metadata dokumen dihapus → {out.name}")


def office_compress(path: str) -> Result:
    """Kompres gambar di dalam file DOCX/XLSX/PPTX."""
    Image = _pil()
    src = Path(path)
    out = unique_path(src, "compressed", src.suffix)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            low = item.filename.lower()
            if "/media/" in low and low.endswith((".jpg", ".jpeg", ".png")):
                try:
                    im = Image.open(io.BytesIO(data))
                    im.load()
                    if max(im.size) > 2000:
                        im.thumbnail((2000, 2000))
                    buf = io.BytesIO()
                    if low.endswith(".png"):
                        if im.mode not in ("RGB", "RGBA"):
                            im = im.convert("RGBA")
                        im.quantize(256, method=2 if im.mode == "RGBA" else 0).save(buf, "PNG", optimize=True)
                    else:
                        _flatten(im).save(buf, "JPEG", quality=70, optimize=True)
                    if buf.tell() < len(data):
                        data = buf.getvalue()
                except Exception:
                    pass
            zout.writestr(item, data, compress_type=zipfile.ZIP_DEFLATED)
    if out.stat().st_size >= src.stat().st_size:
        out.unlink(missing_ok=True)
        return Result([], f"{src.name} sudah teroptimasi, tidak bisa dikompres lebih kecil.")
    return Result([str(out)], f"{out.name}\n{_size_report(src, out)}")


# --------------------------------------------------------------------------- #
#  Teks / Markdown / HTML
# --------------------------------------------------------------------------- #

def text_to_pdf(path: str) -> Result:
    src = Path(path)
    ext = src.suffix.lower()
    raw = _read_text(src)
    if ext in (".html", ".htm"):
        body = raw
    elif ext in (".md", ".markdown"):
        body = _text_to_html(raw, markdown=True)
    else:
        body = f"<pre style='white-space: pre-wrap;'>{html.escape(raw)}</pre>" \
            if ext in (".log", ".json", ".xml", ".csv") else _text_to_html(raw)
    out = unique_path(src, "converted", ".pdf")
    html_to_pdf(body, out, base_dir=str(src.parent))
    return Result([str(out)], f"{src.name} → {out.name}")


def text_to_docx(path: str) -> Result:
    from docx import Document
    src = Path(path)
    raw = _read_text(src)
    doc = Document()
    if src.suffix.lower() in (".md", ".markdown"):
        for line in raw.splitlines():
            s = line.rstrip()
            if s.startswith("#"):
                level = min(len(s) - len(s.lstrip("#")), 4)
                doc.add_heading(s.lstrip("#").strip(), level=level)
            elif s.lstrip().startswith(("- ", "* ")):
                doc.add_paragraph(s.lstrip()[2:], style="List Bullet")
            else:
                doc.add_paragraph(s)
    else:
        for para in raw.split("\n"):
            doc.add_paragraph(para)
    out = unique_path(src, "converted", ".docx")
    doc.save(out)
    return Result([str(out)], f"{src.name} → {out.name}")


def text_to_epub(path: str) -> Result:
    src = Path(path)
    raw = _read_text(src)
    body = _text_to_html(raw, markdown=src.suffix.lower() in (".md", ".markdown"))
    out = unique_path(src, "converted", ".epub")
    build_epub(out, src.stem, "Tidak diketahui", [(src.stem, _xhtml_safe(body))])
    return Result([str(out)], f"{src.name} → {out.name}")


def text_to_html(path: str) -> Result:
    src = Path(path)
    raw = _read_text(src)
    body = _text_to_html(raw, markdown=src.suffix.lower() in (".md", ".markdown"))
    page = (f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{html.escape(src.stem)}</title>"
            f"<style>{_CSS} body{{max-width:820px;margin:40px auto;}}</style></head><body>{body}</body></html>")
    out = unique_path(src, "converted", ".html")
    out.write_text(page, encoding="utf-8")
    return Result([str(out)], f"{src.name} → {out.name}")


# --------------------------------------------------------------------------- #
#  Umum
# --------------------------------------------------------------------------- #

def make_zip(paths: list[str]) -> Result:
    first = Path(paths[0])
    out = unique_path(first, "compressed" if len(paths) == 1 else "archive", ".zip")
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for p in map(Path, paths):
            if p.is_dir():
                for f in p.rglob("*"):
                    if f.is_file():
                        z.write(f, Path(p.name) / f.relative_to(p))
            elif p.is_file():
                z.write(p, p.name)
    return Result([str(out)], f"{len(paths)} item → {out.name} ({human_size(out.stat().st_size)})")


def extract_zip(path: str) -> Result:
    src = Path(path)
    dest = src.parent / f"{src.stem}_extracted"
    i = 1
    while dest.exists():
        dest = src.parent / f"{src.stem}_extracted ({i})"
        i += 1
    with zipfile.ZipFile(src) as z:
        # cegah path traversal
        root = dest.resolve()
        for m in z.infolist():
            target = (dest / m.filename).resolve()
            if not str(target).startswith(str(root)):
                raise ProcessingError("Arsip berisi path yang tidak aman.")
        z.extractall(dest)
    return Result([str(dest)], f"Diekstrak ke {dest.name}")


# --------------------------------------------------------------------------- #
#  EPUB builder (tanpa dependensi eksternal)
# --------------------------------------------------------------------------- #

def _xhtml_safe(fragment: str) -> str:
    """Perbaiki tag void agar valid XHTML."""
    import re
    fragment = re.sub(r"<(br|hr)\s*>", r"<\1/>", fragment)
    fragment = fragment.replace("&nbsp;", "&#160;")
    return fragment


def build_epub(out: Path, title: str, author: str, chapters: list[tuple[str, str]]):
    uid = f"urn:uuid:{uuid.uuid4()}"
    t = html.escape(title)
    a = html.escape(author)
    with zipfile.ZipFile(out, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0" encoding="UTF-8"?>'
                   '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                   '<rootfiles><rootfile full-path="OEBPS/content.opf" '
                   'media-type="application/oebps-package+xml"/></rootfiles></container>',
                   compress_type=zipfile.ZIP_DEFLATED)
        manifest, spine, navpoints, navlis = [], [], [], []
        for i, (ctitle, body) in enumerate(chapters, 1):
            fname = f"chap{i:04d}.xhtml"
            ct = html.escape(ctitle)
            z.writestr(f"OEBPS/{fname}",
                       '<?xml version="1.0" encoding="utf-8"?><!DOCTYPE html>'
                       '<html xmlns="http://www.w3.org/1999/xhtml"><head>'
                       f'<title>{ct}</title><meta charset="utf-8"/></head><body>'
                       f'<h2>{ct}</h2>{body}</body></html>', compress_type=zipfile.ZIP_DEFLATED)
            manifest.append(f'<item id="c{i}" href="{fname}" media-type="application/xhtml+xml"/>')
            spine.append(f'<itemref idref="c{i}"/>')
            navpoints.append(f'<navPoint id="n{i}" playOrder="{i}"><navLabel><text>{ct}</text>'
                             f'</navLabel><content src="{fname}"/></navPoint>')
            navlis.append(f'<li><a href="{fname}">{ct}</a></li>')
        z.writestr("OEBPS/nav.xhtml",
                   '<?xml version="1.0" encoding="utf-8"?><!DOCTYPE html>'
                   '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">'
                   f'<head><title>{t}</title></head><body><nav epub:type="toc"><h1>Daftar Isi</h1><ol>'
                   + "".join(navlis) + '</ol></nav></body></html>', compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("OEBPS/toc.ncx",
                   '<?xml version="1.0" encoding="UTF-8"?>'
                   '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">'
                   f'<head><meta name="dtb:uid" content="{uid}"/></head>'
                   f'<docTitle><text>{t}</text></docTitle><navMap>' + "".join(navpoints)
                   + '</navMap></ncx>', compress_type=zipfile.ZIP_DEFLATED)
        modified = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")
        z.writestr("OEBPS/content.opf",
                   '<?xml version="1.0" encoding="UTF-8"?>'
                   '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">'
                   '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
                   f'<dc:identifier id="bookid">{uid}</dc:identifier><dc:title>{t}</dc:title>'
                   f'<dc:creator>{a}</dc:creator><dc:language>id</dc:language>'
                   f'<meta property="dcterms:modified">{modified}</meta></metadata>'
                   '<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
                   '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>'
                   + "".join(manifest) + '</manifest><spine toc="ncx">' + "".join(spine)
                   + '</spine></package>', compress_type=zipfile.ZIP_DEFLATED)
